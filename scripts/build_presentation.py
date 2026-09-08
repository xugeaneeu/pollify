#!/usr/bin/env python3
"""Generate the 5-minute database-focused presentation for Pollify.

Run from the repo root:
    python3 scripts/build_presentation.py

Produces both English and Russian decks under docs/.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT_DIR = Path(__file__).resolve().parent.parent / "docs"

# Palette (matches the SPA's index.css)
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
INDIGO_DARK = RGBColor(0x36, 0x2D, 0xA8)
INDIGO_SOFT = RGBColor(0xE0, 0xE7, 0xFF)
TEXT = RGBColor(0x11, 0x18, 0x27)
MUTED = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT_AMBER = RGBColor(0x92, 0x40, 0x0E)

prs_width = Inches(13.333)


# ── i18n strings ───────────────────────────────────────────────────────


def t(locale):
    if locale == "ru":
        return RU
    return EN


EN = {
    "title_eyebrow": "POLLIFY",
    "title_main": "Voting & quorum-moderated polls",
    "title_sub": "Database design through the lens of theory",
    "title_stack": "Postgres 17 · 7 tables · 2 enum types · 3 triggers · 9 indexes",
    "title_author": "Daniil Puchkov",
    "footer_brand": "Pollify · database design",

    "s2_eyebrow": "01 · Schema at a glance",
    "s2_title":   "Seven tables that encode the whole domain",
    "s2_bullets": [
        "users — identity + role (USER / ADMIN)",
        "polls — title, question, schedule, settings flags",
        "options — fixed choices for option-based polls",
        "poll_participants — who voted (PK = poll_id, user_id)",
        "votes — what was chosen (option_id XOR custom_text)",
        "reports — moderation tickets opened by users",
        "report_reviews — admin decisions (PK = report_id, admin_id)",
    ],
    "s2_inv_title": "POSTGRES 17 — INVENTORY",
    "s2_inv_rows": [
        ("Tables",            "7"),
        ("Enum types",        "2 — report_status, review_decision"),
        ("CHECK constraints", "5 (poll schedule, max_choices, "
                              "mixed-mode, vote payload, role)"),
        ("Composite PKs",     "2 (poll_participants, report_reviews)"),
        ("Foreign keys",      "11 with ON DELETE CASCADE / SET NULL"),
        ("Triggers",          "3 (admin-only review, no self-review, "
                              "vote-option matches poll)"),
        ("Indexes",           "9 (FK lookups + visibility filter)"),
    ],

    "s3_eyebrow": "02 · the two-table trick",
    "s3_title":   "Anonymity by structure, not by encryption",
    "s3_goal":    "Goal: enforce \"one user, one vote\" without storing "
                  "the user→choice link for anonymous polls.",
    "s3_left_h":  "poll_participants",
    "s3_left_p":  "WHO took part — always recorded.",
    "s3_left_c":  "PRIMARY KEY (poll_id, user_id)\nvoted_at TIMESTAMPTZ",
    "s3_right_h": "votes",
    "s3_right_p": "WHAT was chosen — user_id NULL when anonymous.",
    "s3_right_c": "user_id UUID NULL\noption_id UUID NULL\ncustom_text TEXT NULL",
    "s3_theory_h":"DATABASE THEORY",
    "s3_theory":  ["Identity ⊥ choice: by separating the two relations, "
                   "the uniqueness invariant is enforced declaratively (PK)",
                   "while the user→vote linkage is severed at the schema "
                   "level — anonymity is a structural property, not a "
                   "runtime guard."],

    "s4_eyebrow": "03 · declarative integrity",
    "s4_title":   "Five CHECKs and two composite PKs replace dozens of guards",
    "s4_code":    ("CHECK (end_at > start_at)\n"
                   "CHECK (NOT is_multiple_choice OR max_choices > 1)\n"
                   "CHECK (NOT (is_multiple_choice AND allow_custom_answer))\n"
                   "CHECK (\n"
                   "    (option_id IS NOT NULL AND custom_text IS NULL) OR\n"
                   "    (option_id IS NULL     AND custom_text IS NOT NULL)\n"
                   ")\n"
                   "CHECK (role IN ('USER', 'ADMIN'))\n\n"
                   "PRIMARY KEY (poll_id, user_id)        -- poll_participants\n"
                   "PRIMARY KEY (report_id, admin_id)     -- report_reviews"),
    "s4_theory_h":"CODD'S RULE 10",
    "s4_theory":  ["\"Integrity constraints must be definable and storable "
                   "in the catalog, not in application programs.\"",
                   "",
                   "Every rule above eliminates a class of bugs at the data "
                   "boundary — a misbehaving handler, a hand-written script, "
                   "or a future microservice cannot bypass them."],

    "s5_eyebrow": "04 · triggers as second-line defense",
    "s5_title":   "Row-level rules that CHECK alone cannot express",
    "s5_bullets": [
        "report_reviews_admin_only — every INSERT/UPDATE asserts the "
        "actor's role is ADMIN.",
        "report_reviews_no_self_review — admin cannot review a report "
        "they themselves filed.",
        "votes_option_matches_poll — a vote's option_id must belong to "
        "the same poll_id (cross-table referential check).",
    ],
    "s5_theory_h":"DEFENSE IN DEPTH",
    "s5_theory":  ["App layer validates first — quick feedback, "
                   "structured error envelopes.",
                   "",
                   "Triggers are the last line — even with direct DB "
                   "access, the invariants hold. Procedural counterpart "
                   "to declarative CHECK rules: row-level computations "
                   "across tables that CHECK can't see."],
    "s5_code":    ("CREATE TRIGGER report_reviews_no_self_review\n"
                   "BEFORE INSERT OR UPDATE ON report_reviews\n"
                   "FOR EACH ROW EXECUTE FUNCTION ensure_report_review_not_self();"),

    "s6_eyebrow": "05 · ACID and quorum",
    "s6_title":   "Vote and review as atomic transactions",
    "s6_left_h":  "VOTE SUBMISSION",
    "s6_left_p":  "Single transaction across two tables",
    "s6_left_c":  ("BEGIN;\n"
                   "INSERT INTO poll_participants\n"
                   "   (poll_id, user_id, voted_at)\n"
                   "VALUES ($1, $2, now());     -- 23505 ⇒ 409\n\n"
                   "INSERT INTO votes (...)     -- N rows for\n"
                   "VALUES (...);               -- multi-choice\n"
                   "COMMIT;"),
    "s6_right_h": "ADMIN REVIEW",
    "s6_right_p": "Review row + report row in lockstep",
    "s6_right_c": ("BEGIN;\n"
                   "INSERT INTO report_reviews\n"
                   "   (report_id, admin_id, decision)\n"
                   "VALUES ($1, $2, $3);   -- triggers run\n\n"
                   "UPDATE reports SET\n"
                   "   status, approval_count, rejection_count\n"
                   "WHERE id = $1;\n"
                   "COMMIT;"),
    "s6_banner":  "ACID + UNIQUE = no app-level locking. "
                  "Re-vote, double-review, race condition? The constraint "
                  "fires first; we map 23505 (unique violation) to a "
                  "clean 409 Conflict.",

    "s7_eyebrow":  "06 · recap",
    "s7_title":    "What Postgres gave us, what stayed in the app",
    "s7_db_h":     "DELEGATED TO POSTGRES",
    "s7_db":       [
        "Uniqueness — composite PKs prevent double votes, double reviews",
        "Referential integrity — FKs cascade deletes, ON DELETE SET NULL",
        "Domain typing — enums for report_status, review_decision",
        "Cross-row invariants — triggers for admin role, self-review, "
        "vote/poll alignment",
        "Atomicity — vote = participation + N votes in one transaction",
    ],
    "s7_app_h":    "STILL IN THE APP LAYER",
    "s7_app":      [
        "Computed status (scheduled / active / completed / hidden) — "
        "depends on now()",
        "Quorum threshold — pure function over the reviews list, parameterised",
        "Anonymity write-strategy — set votes.user_id = NULL when "
        "poll.is_anonymous",
        "Aggregations — option counts, percentages, voter listings",
        "Caching, paginated counts, index-friendly query shapes",
    ],
    "s7_banner":   "Schema as a contract: every business rule we encoded "
                   "once, in one migration, lives forever.",

    "notes": [
        "Hi, I'm Daniil. Today I'll walk you through Pollify — a polling "
        "and voting service. The focus is on how we leaned on Postgres "
        "to encode the domain rules, and how those choices line up with "
        "database theory.",

        "The whole domain fits into seven tables. Users, polls, and "
        "options for fixed choices. The interesting pair is "
        "poll_participants and votes — we keep the fact someone "
        "participated separate from what they actually chose. Reports "
        "and report_reviews drive moderation. Two enum types for report "
        "status and review decision, three triggers, nine indexes — all "
        "in a single migration file.",

        "Here's the headline design choice. We need two invariants: "
        "one user, one vote; and for anonymous polls, no link between "
        "user and choice. Solving both at the schema level is what this "
        "two-table split does. poll_participants has a composite primary "
        "key — poll_id and user_id together — so duplicate participation "
        "is impossible. Votes carries the actual answer, but its user_id "
        "is nullable. For anonymous polls we leave user_id NULL. The "
        "mapping from user to vote literally does not exist in the row. "
        "This is the database analog of secret-ballot voting — separation "
        "of identity from data — enforced by structure, not by application "
        "checks. Even direct SQL access can't reconstruct the link.",

        "Five CHECK constraints and two composite primary keys do most "
        "of the work. End time after start time. Multi-choice means "
        "max_choices ≥ 2. Multi-choice and free-text are mutually "
        "exclusive. A vote row carries an option XOR a custom text — "
        "never both, never neither. Role constrained to USER or ADMIN. "
        "This maps onto Codd's tenth rule: integrity constraints belong "
        "in the catalog, not in application programs. Each rule closes "
        "off a class of bugs at the data boundary — a misbehaving handler "
        "or hand-written script can't bypass them.",

        "Some invariants need procedural row-level logic that CHECK can't "
        "express. Three triggers. Only admins can insert report_reviews. "
        "Admins can't review their own reports — that one even joins "
        "reports and report_reviews. A vote's option must belong to the "
        "same poll as the vote — referential integrity across tables. "
        "Application code validates first for fast feedback; triggers "
        "are the last line of defense — they fire even on direct DB access.",

        "Vote submission and admin review are both atomic transactions. "
        "Vote submission inserts into poll_participants and votes in one "
        "transaction. If a user tries to vote twice, the UNIQUE violation "
        "on poll_participants fires — we map error code 23505 to a clean "
        "409 Conflict, and the votes inserts roll back. Admin review "
        "inserts into report_reviews and updates the report row in one "
        "transaction; the trigger checks role and self-review before "
        "either succeeds. The quorum logic itself is a pure function over "
        "the reviews list, run by the application right before the UPDATE. "
        "ACID plus UNIQUE replaces application-level locking.",

        "Recap. What Postgres gave us: uniqueness, referential integrity, "
        "domain typing via enums, cross-row invariants via triggers, "
        "atomicity across multi-table writes. What stayed in the "
        "application: anything that depends on the current time — "
        "computing scheduled / active / completed / hidden. The quorum "
        "threshold itself, parameterized. Aggregations and pagination. "
        "Schema as a contract: every business rule we encoded once, in "
        "one migration, lives forever — outliving any future code that "
        "touches the database. Thank you.",
    ],
}


RU = {
    "title_eyebrow": "POLLIFY",
    "title_main":    "Голосования с кворумной модерацией",
    "title_sub":     "Проектирование БД через призму теории",
    "title_stack":   "Postgres 17 · 7 таблиц · 2 ENUM-типа · 3 триггера · 9 индексов",
    "title_author":  "Даниил Пучков",
    "footer_brand":  "Pollify · проектирование БД",

    "s2_eyebrow": "01 · ОБЗОР СХЕМЫ",
    "s2_title":   "Семь таблиц описывают всю предметную область",
    "s2_bullets": [
        "users — идентичность + роль (USER / ADMIN)",
        "polls — заголовок, вопрос, расписание, флаги настроек",
        "options — варианты для опросов с фиксированным выбором",
        "poll_participants — кто голосовал (PK = poll_id, user_id)",
        "votes — что выбрано (option_id XOR custom_text)",
        "reports — жалобы, поданные пользователями",
        "report_reviews — решения админов (PK = report_id, admin_id)",
    ],
    "s2_inv_title": "POSTGRES 17 — СОСТАВ",
    "s2_inv_rows": [
        ("Таблицы",            "7"),
        ("ENUM-типы",          "2 — report_status, review_decision"),
        ("CHECK-ограничения",  "5 (расписание, max_choices, "
                               "смешанный режим, payload голоса, роль)"),
        ("Составные PK",       "2 (poll_participants, report_reviews)"),
        ("Внешние ключи",      "11 с ON DELETE CASCADE / SET NULL"),
        ("Триггеры",           "3 (review только для админа, запрет "
                               "self-review, option принадлежит poll)"),
        ("Индексы",            "9 (FK + фильтр видимости)"),
    ],

    "s3_eyebrow": "02 · ПРИЁМ ДВУХ ТАБЛИЦ",
    "s3_title":   "Анонимность через структуру, а не шифрованием",
    "s3_goal":    "Цель: гарантировать «один пользователь — один голос», "
                  "не сохраняя связь пользователь→выбор для анонимных опросов.",
    "s3_left_h":  "poll_participants",
    "s3_left_p":  "КТО участвовал — всегда фиксируется.",
    "s3_left_c":  "PRIMARY KEY (poll_id, user_id)\nvoted_at TIMESTAMPTZ",
    "s3_right_h": "votes",
    "s3_right_p": "ЧТО выбрано — user_id NULL для анонимных.",
    "s3_right_c": "user_id UUID NULL\noption_id UUID NULL\ncustom_text TEXT NULL",
    "s3_theory_h":"ТЕОРИЯ БАЗ ДАННЫХ",
    "s3_theory":  ["Идентичность ⊥ выбор: разделяя два отношения, "
                   "инвариант уникальности обеспечивается декларативно (PK),",
                   "а связь пользователь→голос разрывается на уровне схемы — "
                   "анонимность становится структурным свойством, "
                   "а не runtime-проверкой."],

    "s4_eyebrow": "03 · ДЕКЛАРАТИВНАЯ ЦЕЛОСТНОСТЬ",
    "s4_title":   "Пять CHECK и два составных PK заменяют десятки проверок",
    "s4_code":    ("CHECK (end_at > start_at)\n"
                   "CHECK (NOT is_multiple_choice OR max_choices > 1)\n"
                   "CHECK (NOT (is_multiple_choice AND allow_custom_answer))\n"
                   "CHECK (\n"
                   "    (option_id IS NOT NULL AND custom_text IS NULL) OR\n"
                   "    (option_id IS NULL     AND custom_text IS NOT NULL)\n"
                   ")\n"
                   "CHECK (role IN ('USER', 'ADMIN'))\n\n"
                   "PRIMARY KEY (poll_id, user_id)        -- poll_participants\n"
                   "PRIMARY KEY (report_id, admin_id)     -- report_reviews"),
    "s4_theory_h":"ПРАВИЛО 10 КОДДА",
    "s4_theory":  ["«Ограничения целостности должны определяться и "
                   "храниться в каталоге, а не в прикладных программах.»",
                   "",
                   "Каждое правило закрывает класс ошибок на границе "
                   "данных: некорректный обработчик, ручной скрипт или "
                   "будущий микросервис не смогут его обойти."],

    "s5_eyebrow": "04 · ТРИГГЕРЫ — ВТОРАЯ ЛИНИЯ ОБОРОНЫ",
    "s5_title":   "Правила уровня строки, которые не выразить через CHECK",
    "s5_bullets": [
        "report_reviews_admin_only — каждый INSERT/UPDATE проверяет, "
        "что роль актора = ADMIN.",
        "report_reviews_no_self_review — админ не может рассматривать "
        "свою же жалобу.",
        "votes_option_matches_poll — option_id голоса должен принадлежать "
        "тому же poll_id (межтабличная ссылочная проверка).",
    ],
    "s5_theory_h":"ОБОРОНА В ГЛУБИНУ",
    "s5_theory":  ["Слой приложения валидирует первым — быстрая обратная "
                   "связь, структурированные ошибки.",
                   "",
                   "Триггеры — последняя линия: даже при прямом доступе "
                   "к БД инварианты соблюдаются. Это процедурный аналог "
                   "декларативных CHECK: проверки уровня строки между "
                   "таблицами, недоступные CHECK."],
    "s5_code":    ("CREATE TRIGGER report_reviews_no_self_review\n"
                   "BEFORE INSERT OR UPDATE ON report_reviews\n"
                   "FOR EACH ROW EXECUTE FUNCTION ensure_report_review_not_self();"),

    "s6_eyebrow": "05 · ACID И КВОРУМ",
    "s6_title":   "Голос и review как атомарные транзакции",
    "s6_left_h":  "ОТПРАВКА ГОЛОСА",
    "s6_left_p":  "Одна транзакция на две таблицы",
    "s6_left_c":  ("BEGIN;\n"
                   "INSERT INTO poll_participants\n"
                   "   (poll_id, user_id, voted_at)\n"
                   "VALUES ($1, $2, now());     -- 23505 ⇒ 409\n\n"
                   "INSERT INTO votes (...)     -- N строк при\n"
                   "VALUES (...);               -- multi-choice\n"
                   "COMMIT;"),
    "s6_right_h": "REVIEW АДМИНА",
    "s6_right_p": "Строка review и строка report — синхронно",
    "s6_right_c": ("BEGIN;\n"
                   "INSERT INTO report_reviews\n"
                   "   (report_id, admin_id, decision)\n"
                   "VALUES ($1, $2, $3);   -- сработают триггеры\n\n"
                   "UPDATE reports SET\n"
                   "   status, approval_count, rejection_count\n"
                   "WHERE id = $1;\n"
                   "COMMIT;"),
    "s6_banner":  "ACID + UNIQUE = без блокировок на уровне приложения. "
                  "Повторный голос, дубль review, гонка? Ограничение "
                  "сработает первым; код 23505 (unique violation) → "
                  "чистый 409 Conflict.",

    "s7_eyebrow":  "06 · ИТОГ",
    "s7_title":    "Что взял на себя Postgres, что осталось в приложении",
    "s7_db_h":     "ОТДЕЛЕГИРОВАНО POSTGRES",
    "s7_db":       [
        "Уникальность — составные PK не дают повторных голосов и review",
        "Ссылочная целостность — FK с CASCADE и ON DELETE SET NULL",
        "Доменные типы — enum для report_status, review_decision",
        "Межстрочные инварианты — триггеры на роль ADMIN, self-review, "
        "привязку option к poll",
        "Атомарность — голос = participation + N votes в одной транзакции",
    ],
    "s7_app_h":    "ОСТАЛОСЬ В ПРИЛОЖЕНИИ",
    "s7_app":      [
        "Вычисляемый статус (scheduled / active / completed / hidden) — "
        "зависит от now()",
        "Порог кворума — чистая функция над списком reviews, параметризованная",
        "Стратегия записи для анонимности — votes.user_id = NULL "
        "при poll.is_anonymous",
        "Агрегации — подсчёт по вариантам, проценты, списки голосовавших",
        "Кэширование, постраничный подсчёт, индекс-дружественные запросы",
    ],
    "s7_banner":   "Схема как контракт: каждое бизнес-правило, "
                   "записанное однажды в одной миграции, живёт вечно.",

    "notes": [
        "Здравствуйте, меня зовут Даниил. Сегодня я расскажу о Pollify — "
        "сервисе для опросов и голосований. Основной фокус — как мы "
        "опираемся на Postgres для кодирования доменных правил, и как "
        "эти решения соотносятся с теорией баз данных.",

        "Вся предметная область укладывается в семь таблиц. Users, "
        "polls и options для фиксированных вариантов. Интересная пара — "
        "poll_participants и votes: мы храним факт участия отдельно от "
        "самого выбора. Reports и report_reviews отвечают за модерацию. "
        "Два enum-типа для статуса жалобы и решения review, три "
        "триггера, девять индексов — всё в одном файле миграции.",

        "Главное архитектурное решение. Нам нужны два инварианта: один "
        "пользователь — один голос; и для анонимных опросов — никакой "
        "связи между пользователем и выбором. Это разделение на две "
        "таблицы решает обе задачи на уровне схемы. У poll_participants "
        "составной первичный ключ — poll_id и user_id вместе — повторное "
        "участие физически невозможно. Votes хранит сам ответ, но "
        "user_id там nullable. Для анонимных опросов оставляем user_id "
        "пустым. Связи между пользователем и голосом просто нет в строке. "
        "Это аналог тайного голосования в БД — разделение идентичности "
        "и данных — обеспеченное структурой, а не runtime-проверками. "
        "Даже прямой SQL не восстановит связь.",

        "Пять CHECK-ограничений и два составных PK делают почти всю "
        "работу. End_at позже start_at. Multi-choice требует "
        "max_choices ≥ 2. Multi-choice и свободный текст взаимоисключают "
        "друг друга. Строка votes несёт option XOR custom_text — никогда "
        "оба, никогда ни одного. Роль ограничена USER или ADMIN. Это "
        "правило 10 Кодда: ограничения целостности должны храниться в "
        "каталоге, а не в прикладных программах. Каждое правило "
        "закрывает класс ошибок на границе данных — некорректный "
        "обработчик или ручной скрипт не смогут его обойти.",

        "Некоторые инварианты требуют процедурной логики уровня строки, "
        "которую CHECK выразить не может. Три триггера. Только админы "
        "могут вставлять строки в report_reviews. Админ не может "
        "рассматривать собственную жалобу — этот триггер даже джойнит "
        "reports с report_reviews. Option голоса должен принадлежать "
        "тому же poll, что и сам голос — ссылочная целостность между "
        "таблицами. Приложение валидирует первым для быстрой обратной "
        "связи; триггеры — последняя линия, они срабатывают даже при "
        "прямом доступе к БД.",

        "Отправка голоса и review админа — обе атомарные транзакции. "
        "Отправка голоса вставляет в poll_participants и votes в одной "
        "транзакции. Если пользователь пытается проголосовать дважды — "
        "срабатывает UNIQUE-нарушение на poll_participants, мы маппим "
        "код 23505 в чистый 409 Conflict, а вставки в votes откатываются. "
        "Review админа вставляет в report_reviews и обновляет строку "
        "report в одной транзакции; триггер проверяет роль и self-review "
        "до того, как что-либо успешно выполнится. Логика кворума — "
        "чистая функция над списком reviews, выполняемая приложением "
        "прямо перед UPDATE. ACID плюс UNIQUE заменяет блокировки на "
        "уровне приложения.",

        "Итог. Что взял на себя Postgres: уникальность, ссылочная "
        "целостность, доменная типизация через enum, межстрочные "
        "инварианты через триггеры, атомарность многотабличных записей. "
        "Что осталось в приложении: всё, что зависит от текущего "
        "времени — вычисление scheduled / active / completed / hidden. "
        "Сам порог кворума, параметризованный. Агрегации и пагинация. "
        "Схема как контракт: каждое бизнес-правило, записанное один раз "
        "в одной миграции, живёт вечно — переживёт любой код, который "
        "будет работать с этой базой. Спасибо.",
    ],
}


# ── helpers ────────────────────────────────────────────────────────────


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def no_line(shape):
    shape.line.fill.background()


def add_rect(slide, x, y, w, h, color):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    fill(rect, color)
    no_line(rect)
    return rect


def add_text(slide, x, y, w, h, text, *,
             size=18, bold=False, color=TEXT, align=PP_ALIGN.LEFT,
             font="Calibri", line_spacing=1.15):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.word_wrap = True
    if isinstance(text, str):
        text = [text]
    for i, line in enumerate(text):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_bullets(slide, x, y, w, h, bullets, *,
                size=16, color=TEXT, font="Calibri",
                line_spacing=1.35):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_before = Pt(0 if i == 0 else 6)
        run = p.add_run()
        run.text = "▸  " + bullet
        run.font.name = font
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return box


def add_code(slide, x, y, w, h, code, *, size=12, color=TEXT, bg=INDIGO_SOFT):
    rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    fill(rect, bg)
    no_line(rect)
    rect.adjustments[0] = 0.06

    box = slide.shapes.add_textbox(x + Emu(160000), y + Emu(140000),
                                   w - Emu(320000), h - Emu(280000))
    tf = box.text_frame
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.word_wrap = True
    lines = code.splitlines()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.15
        run = p.add_run()
        run.text = line if line else " "
        run.font.name = "Consolas"
        run.font.size = Pt(size)
        run.font.color.rgb = color


def slide_title(slide, eyebrow, title):
    add_rect(slide, 0, 0, prs_width, Inches(0.06), INDIGO)
    if eyebrow:
        add_text(slide, Inches(0.7), Inches(0.45), Inches(11), Inches(0.4),
                 eyebrow, size=11, bold=True, color=INDIGO)
    add_text(slide, Inches(0.7), Inches(0.8), Inches(12), Inches(1.0),
             title, size=32, bold=True, color=TEXT)


def add_footer(slide, page, total, brand):
    add_text(slide, Inches(0.7), Inches(7.05), Inches(8), Inches(0.3),
             brand, size=10, color=MUTED)
    add_text(slide, Inches(11.6), Inches(7.05), Inches(1.4), Inches(0.3),
             f"{page} / {total}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def set_notes(slide, text):
    notes = slide.notes_slide.notes_text_frame
    notes.text = text


# ── slides ─────────────────────────────────────────────────────────────


def title_slide(slide, L):
    add_rect(slide, 0, 0, prs_width, Inches(7.5), INDIGO)
    add_rect(slide, Inches(0.7), Inches(2.4), Inches(0.18), Inches(2.4), WHITE)

    add_text(slide, Inches(1.2), Inches(2.3), Inches(11), Inches(0.5),
             L["title_eyebrow"], size=14, bold=True, color=WHITE)
    add_text(slide, Inches(1.2), Inches(2.7), Inches(11), Inches(2.0),
             L["title_main"], size=46, bold=True, color=WHITE)
    add_text(slide, Inches(1.2), Inches(3.9), Inches(11), Inches(1.2),
             L["title_sub"], size=24, color=RGBColor(0xC7, 0xD2, 0xFE))
    add_text(slide, Inches(1.2), Inches(6.4), Inches(11), Inches(0.4),
             L["title_stack"], size=14, color=RGBColor(0xC7, 0xD2, 0xFE))
    add_text(slide, Inches(1.2), Inches(6.75), Inches(11), Inches(0.4),
             L["title_author"], size=14, color=WHITE)


def schema_overview_slide(slide, L):
    slide_title(slide, L["s2_eyebrow"], L["s2_title"])
    add_bullets(slide, Inches(0.7), Inches(2.0), Inches(6.3), Inches(4.6),
                L["s2_bullets"], size=16)

    add_rect(slide, Inches(7.6), Inches(2.0), Inches(5.0), Inches(4.6), WHITE)
    add_text(slide, Inches(7.9), Inches(2.2), Inches(4.4), Inches(0.4),
             L["s2_inv_title"], size=11, bold=True, color=INDIGO)

    y = Inches(2.7)
    for label, value in L["s2_inv_rows"]:
        add_text(slide, Inches(7.9), y, Inches(2.0), Inches(0.4),
                 label, size=12, color=MUTED, bold=True)
        add_text(slide, Inches(9.9), y, Inches(2.6), Inches(0.4),
                 value, size=12, color=TEXT)
        y += Inches(0.5)


def anonymity_slide(slide, L):
    slide_title(slide, L["s3_eyebrow"], L["s3_title"])
    add_text(slide, Inches(0.7), Inches(1.95), Inches(12), Inches(0.5),
             L["s3_goal"], size=14, color=MUTED)

    add_rect(slide, Inches(0.7), Inches(2.7), Inches(5.6), Inches(2.0), WHITE)
    add_text(slide, Inches(0.95), Inches(2.85), Inches(5), Inches(0.4),
             L["s3_left_h"], size=14, bold=True, color=INDIGO_DARK)
    add_text(slide, Inches(0.95), Inches(3.2), Inches(5), Inches(0.5),
             L["s3_left_p"], size=12, color=MUTED)
    add_code(slide, Inches(0.95), Inches(3.55), Inches(5.0), Inches(1.05),
             L["s3_left_c"], size=12)

    add_rect(slide, Inches(7.0), Inches(2.7), Inches(5.6), Inches(2.0), WHITE)
    add_text(slide, Inches(7.25), Inches(2.85), Inches(5), Inches(0.4),
             L["s3_right_h"], size=14, bold=True, color=INDIGO_DARK)
    add_text(slide, Inches(7.25), Inches(3.2), Inches(5), Inches(0.5),
             L["s3_right_p"], size=12, color=MUTED)
    add_code(slide, Inches(7.25), Inches(3.55), Inches(5.0), Inches(1.05),
             L["s3_right_c"], size=12)

    add_rect(slide, Inches(0.7), Inches(5.0), Inches(11.9), Inches(1.6),
             INDIGO_SOFT)
    add_text(slide, Inches(0.95), Inches(5.15), Inches(12), Inches(0.4),
             L["s3_theory_h"], size=11, bold=True, color=INDIGO_DARK)
    add_text(slide, Inches(0.95), Inches(5.5), Inches(11.6), Inches(1.0),
             L["s3_theory"], size=14, line_spacing=1.3)


def declarative_integrity_slide(slide, L):
    slide_title(slide, L["s4_eyebrow"], L["s4_title"])
    add_code(slide, Inches(0.7), Inches(2.0), Inches(7.6), Inches(4.4),
             L["s4_code"], size=13)

    add_rect(slide, Inches(8.6), Inches(2.0), Inches(4.0), Inches(4.4),
             INDIGO_SOFT)
    add_text(slide, Inches(8.85), Inches(2.15), Inches(3.6), Inches(0.4),
             L["s4_theory_h"], size=11, bold=True, color=INDIGO_DARK)
    add_text(slide, Inches(8.85), Inches(2.5), Inches(3.6), Inches(3.0),
             L["s4_theory"], size=12, color=TEXT, line_spacing=1.35)


def triggers_slide(slide, L):
    slide_title(slide, L["s5_eyebrow"], L["s5_title"])
    add_bullets(slide, Inches(0.7), Inches(2.0), Inches(7.5), Inches(3.0),
                L["s5_bullets"], size=15)

    add_rect(slide, Inches(8.6), Inches(2.0), Inches(4.0), Inches(4.4),
             INDIGO_SOFT)
    add_text(slide, Inches(8.85), Inches(2.15), Inches(3.6), Inches(0.4),
             L["s5_theory_h"], size=11, bold=True, color=INDIGO_DARK)
    add_text(slide, Inches(8.85), Inches(2.5), Inches(3.6), Inches(3.0),
             L["s5_theory"], size=12, color=TEXT, line_spacing=1.35)

    add_code(slide, Inches(0.7), Inches(5.1), Inches(7.5), Inches(1.5),
             L["s5_code"], size=12)


def transactions_slide(slide, L):
    slide_title(slide, L["s6_eyebrow"], L["s6_title"])

    add_rect(slide, Inches(0.7), Inches(2.0), Inches(5.9), Inches(4.0), WHITE)
    add_text(slide, Inches(0.95), Inches(2.15), Inches(5), Inches(0.4),
             L["s6_left_h"], size=11, bold=True, color=INDIGO_DARK)
    add_text(slide, Inches(0.95), Inches(2.5), Inches(5.4), Inches(0.6),
             L["s6_left_p"], size=15, bold=True)
    add_code(slide, Inches(0.95), Inches(3.1), Inches(5.4), Inches(2.7),
             L["s6_left_c"], size=12)

    add_rect(slide, Inches(6.9), Inches(2.0), Inches(5.9), Inches(4.0), WHITE)
    add_text(slide, Inches(7.15), Inches(2.15), Inches(5), Inches(0.4),
             L["s6_right_h"], size=11, bold=True, color=INDIGO_DARK)
    add_text(slide, Inches(7.15), Inches(2.5), Inches(5.4), Inches(0.6),
             L["s6_right_p"], size=15, bold=True)
    add_code(slide, Inches(7.15), Inches(3.1), Inches(5.4), Inches(2.7),
             L["s6_right_c"], size=12)

    add_rect(slide, Inches(0.7), Inches(6.2), Inches(12.1), Inches(0.7),
             INDIGO_SOFT)
    add_text(slide, Inches(0.95), Inches(6.32), Inches(11.7), Inches(0.5),
             L["s6_banner"], size=13, color=TEXT)


def recap_slide(slide, L):
    slide_title(slide, L["s7_eyebrow"], L["s7_title"])

    add_text(slide, Inches(0.7), Inches(2.0), Inches(6), Inches(0.4),
             L["s7_db_h"], size=12, bold=True, color=INDIGO)
    add_bullets(slide, Inches(0.7), Inches(2.4), Inches(6.0), Inches(4.0),
                L["s7_db"], size=14)

    add_text(slide, Inches(7.0), Inches(2.0), Inches(6), Inches(0.4),
             L["s7_app_h"], size=12, bold=True, color=ACCENT_AMBER)
    add_bullets(slide, Inches(7.0), Inches(2.4), Inches(6.0), Inches(4.0),
                L["s7_app"], size=14)

    add_rect(slide, Inches(0.7), Inches(6.2), Inches(12.1), Inches(0.7),
             INDIGO)
    add_text(slide, Inches(0.95), Inches(6.32), Inches(11.7), Inches(0.5),
             L["s7_banner"], size=14, color=WHITE, bold=True)


# ── build ──────────────────────────────────────────────────────────────


def build(locale, out_path):
    L = t(locale)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]
    pages = [
        title_slide,
        schema_overview_slide,
        anonymity_slide,
        declarative_integrity_slide,
        triggers_slide,
        transactions_slide,
        recap_slide,
    ]
    for i, page in enumerate(pages):
        s = prs.slides.add_slide(blank)
        page(s, L)
        set_notes(s, L["notes"][i])

    total = len(prs.slides)
    for i, slide in enumerate(prs.slides):
        if i == 0:
            continue
        add_footer(slide, i + 1, total, L["footer_brand"])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    print(f"wrote {out_path.relative_to(Path.cwd())}")


if __name__ == "__main__":
    build("en", OUT_DIR / "Pollify-Database-Talk.pptx")
    build("ru", OUT_DIR / "Pollify-Database-Talk.ru.pptx")
